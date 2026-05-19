import time
import re
import os
import pandas as pd
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.options import Options
from selenium.common.exceptions import TimeoutException, StaleElementReferenceException

# ============================================================
# CONFIGURAÇÕES
# ============================================================
ANO_LETIVO      = "2026"
URL_BASE        = "https://dedmais.educacao.mg.gov.br/"
ARQUIVO_SAIDA   = f"diario_conteudos_{ANO_LETIVO}.xlsx"



# ============================================================
# SETUP DO NAVEGADOR
# ============================================================
def criar_driver():
    options = Options()
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--disable-extensions")
    options.add_argument("--disable-gpu")
    options.add_argument("--window-size=1400,900")
    driver = webdriver.Chrome(options=options)
    return driver

def wait_for(driver, by, selector, timeout=15):
    return WebDriverWait(driver, timeout).until(
        EC.presence_of_element_located((by, selector))
    )

def completar_data(data_ddmm, ano=ANO_LETIVO):
    partes = data_ddmm.strip().split("/")
    if len(partes) == 2:
        return f"{partes[0]}/{partes[1]}/{ano}"
    elif len(partes) == 3:
        return data_ddmm.strip()
    return data_ddmm.strip()

def parsear_cabecalho(texto):
    """
    Ex: '1º ADMINISTRAÇÃO EM ENSINO REGULAR PROFISSIONAL 1 (2026) -
         Sustentabilidade E Responsabilidade Socioambiental - 2111200 - Manhã -
         Avenida Delvito Alves Da Silva, 888'
    """
    partes = [p.strip() for p in texto.split(" - ")]
    turma      = partes[0] if len(partes) > 0 else ""
    disciplina = partes[1] if len(partes) > 1 else ""
    turno      = partes[3].upper() if len(partes) > 3 else ""
    escola     = partes[4] if len(partes) > 4 else ""
    return turma, disciplina, turno, escola

def extrair_etapa(turma_texto):
    match = re.search(r"(\d+)[oºª°]", turma_texto)
    if match:
        return f"{match.group(1)}º ano"
    return ""

# ============================================================
# CLICA NO BOTÃO "AULAS" (botão laranja v-btn.bg-orange)
# ============================================================
def clicar_botao_aulas(driver):
    """
    Localiza o botão laranja dentro do wrapper que contém o texto 'Aulas'
    e clica nele. Retorna True se conseguiu, False caso contrário.
    Seletor: div.d-flex.flex-column.align-center que contenha 'Aulas'
             → button.v-btn.bg-orange dentro dele.
    """
    try:
        # Aguarda os botões de menu da turma aparecerem
        WebDriverWait(driver, 15).until(
            EC.presence_of_element_located(
                (By.CSS_SELECTOR, "button.v-btn.bg-orange")
            )
        )

        # Busca todos os wrappers d-flex.flex-column.align-center
        wrappers = driver.find_elements(
            By.CSS_SELECTOR, "div.d-flex.flex-column.align-center"
        )

        for wrapper in wrappers:
            if wrapper.text.strip() == "Aulas":
                btn = wrapper.find_element(By.CSS_SELECTOR, "button.v-btn.bg-orange")
                driver.execute_script("arguments[0].click();", btn)
                return True

        # Fallback: clica direto no único botão bg-orange da página
        btn = driver.find_element(By.CSS_SELECTOR, "button.v-btn.bg-orange")
        driver.execute_script("arguments[0].click();", btn)
        return True

    except Exception as e:
        print(f"    [!] Não encontrou botão Aulas: {e}")
        return False

# ============================================================
# COLETA OS DADOS DE AULAS DA PÁGINA DE AULAS
# ============================================================
def coletar_aulas_da_pagina(driver):
    try:
        WebDriverWait(driver, 15).until(
            EC.presence_of_element_located((By.CSS_SELECTOR, ".v-card.my-2"))
        )
    except TimeoutException:
        return [], ""

    time.sleep(0.8)

    # Cabeçalho com info da turma
    cabecalho_els = driver.find_elements(By.CSS_SELECTOR, ".font-weight-bold.my-2")
    cabecalho_texto = cabecalho_els[0].text.strip() if cabecalho_els else ""

    # Cards de aula
    cards = driver.find_elements(By.CSS_SELECTOR, ".v-card.my-2")
    aulas = []
    for card in cards:
        try:
            h3_els       = card.find_elements(By.TAG_NAME, "h3")
            conteudo_els = card.find_elements(By.CSS_SELECTOR, ".dcAula")
            if not h3_els or not conteudo_els:
                continue
            data_raw = h3_els[0].text.strip()
            conteudo = conteudo_els[0].text.strip()
            # Só aceita datas no formato dd/mm
            if re.match(r"^\d{2}/\d{2}$", data_raw) and conteudo:
                aulas.append({
                    "data":     completar_data(data_raw),
                    "conteudo": conteudo
                })
        except StaleElementReferenceException:
            continue

    return aulas, cabecalho_texto

# ============================================================
# FLUXO PRINCIPAL
# ============================================================
def main():
    print("[*] Iniciando coleta do DED...")
    driver = criar_driver()
    todos_dados = []

    try:
        driver.get(URL_BASE)
        time.sleep(3)

        if "login" in driver.current_url.lower() or "auth" in driver.current_url.lower():
            print("[!] Sessão não detectada. Faça login manualmente e pressione Enter...")
            input()
            time.sleep(2)

        print("[*] Aguardando lista de turmas...")
        try:
            wait_for(driver, By.CSS_SELECTOR, ".v-card.v-card--link", timeout=20)
        except TimeoutException:
            print("[!] Página principal não carregou. Verifique o login.")
            driver.quit()
            return

        # Scroll para carregar todos os cards (lazy loading)
        print("[*] Carregando todos os cards...")
        ultima_altura = driver.execute_script("return document.body.scrollHeight")
        for _ in range(10):
            driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
            time.sleep(1)
            nova_altura = driver.execute_script("return document.body.scrollHeight")
            if nova_altura == ultima_altura:
                break
            ultima_altura = nova_altura
        driver.execute_script("window.scrollTo(0, 0);")
        time.sleep(1)

        # Salva metadados dos cards antes de navegar
        cards_info = []
        cards = driver.find_elements(By.CSS_SELECTOR, ".v-card.v-card--link")
        for card in cards:
            try:
                linhas = [l.strip() for l in card.text.split("\n") if l.strip()]
                if len(linhas) >= 4:
                    cards_info.append({
                        "etapa":      linhas[0],
                        "turma":      linhas[1],
                        "escola":     linhas[2],
                        "disciplina": linhas[3],
                        "turno":      linhas[4] if len(linhas) > 4 else ""
                    })
            except Exception:
                continue

        print(f"[*] {len(cards_info)} turmas encontradas. Iniciando coleta...\n")

        for i, meta in enumerate(cards_info):
            print(f"  [{i+1}/{len(cards_info)}] {meta['turma'][:60]}...")

            try:
                # ── 1. Volta para a lista e clica no card ──
                driver.get(URL_BASE)
                time.sleep(2)
                wait_for(driver, By.CSS_SELECTOR, ".v-card.v-card--link", timeout=15)

                cards_atuais = driver.find_elements(By.CSS_SELECTOR, ".v-card.v-card--link")
                if i >= len(cards_atuais):
                    print(f"    [!] Card {i} não encontrado, pulando...")
                    continue

                driver.execute_script("arguments[0].click();", cards_atuais[i])
                time.sleep(2)

                # ── 2. Aguarda a página da turma (Visão Geral) ──
                wait_for(driver, By.CSS_SELECTOR, "button.v-btn.bg-orange", timeout=15)

                # ── 3. Clica no botão laranja "Aulas" ──
                clicou = clicar_botao_aulas(driver)
                if not clicou:
                    print(f"    [!] Não conseguiu clicar em Aulas, pulando...")
                    continue

                time.sleep(2)

                # ── 4. Coleta os dados de aulas ──
                aulas, cabecalho = coletar_aulas_da_pagina(driver)

                if cabecalho:
                    turma_nome, disciplina, turno, escola = parsear_cabecalho(cabecalho)
                    etapa      = extrair_etapa(turma_nome)
                    turma_limpa = re.sub(r"\s*-\s*\d{6,}\s*$", "", turma_nome).strip()
                else:
                    etapa       = meta["etapa"]
                    turma_limpa = meta["turma"]
                    disciplina  = meta["disciplina"]
                    turno       = meta["turno"]
                    escola      = meta["escola"]

                print(f"    -> {len(aulas)} aulas coletadas.")

                for aula in aulas:
                    todos_dados.append({
                        "EtapadaMatrícula":  etapa,
                        "Turma":             turma_limpa,
                        "Escola":            escola,
                        "Disciplina":        disciplina,
                        "Turno":             turno.upper(),
                        "Data":              aula["data"],
                        "ConteudoLecionado": aula["conteudo"],
                    })

            except Exception as e:
                print(f"    [!] Erro inesperado: {e}")

    finally:
        driver.quit()

    # ============================================================
    # GERA A PLANILHA EXCEL
    # ============================================================
    if not todos_dados:
        print("\n[!] Nenhum dado coletado.")
        return

    df = pd.DataFrame(todos_dados, columns=[
        "EtapadaMatrícula", "Turma", "Escola",
        "Disciplina", "Turno", "Data", "ConteudoLecionado"
    ])

    df.drop_duplicates(inplace=True)

    df["_data_ord"] = pd.to_datetime(df["Data"], format="%d/%m/%Y", errors="coerce")
    df.sort_values(["Turma", "Disciplina", "_data_ord"], inplace=True)
    df.drop(columns=["_data_ord"], inplace=True)
    df.reset_index(drop=True, inplace=True)

    with pd.ExcelWriter(ARQUIVO_SAIDA, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Conteúdos Lecionados")

        ws = writer.sheets["Conteúdos Lecionados"]
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        header_fill = PatternFill("solid", fgColor="4B0082")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        for col_idx in range(1, len(df.columns) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22

        larguras = {
            "EtapadaMatrícula":  14,
            "Turma":             58,
            "Escola":            40,
            "Disciplina":        45,
            "Turno":             10,
            "Data":              13,
            "ConteudoLecionado": 85,
        }
        for col_idx, col_name in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = larguras.get(col_name, 20)

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    print(f"\n[✓] Planilha salva: {ARQUIVO_SAIDA}")
    print(f"[✓] Total de registros: {len(df)}")
    print(f"\nPrimeiros registros:\n{df.head(5).to_string(index=False)}")

if __name__ == "__main__":
    main()